"""Exports. CSV is stdlib and always works; XLSX is a bonus when openpyxl is present."""
import csv
import io
from typing import List

COLUMNS = ["id", "status", "mandatory", "category", "field", "requirement",
           "evidence", "page", "verified", "quote"]


def rows_to_csv(rows: List[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8")


def rows_to_xlsx(rows: List[dict], decision: str = "") -> bytes | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return None
    fills = {"Met": "C6EFCE", "Not Met": "FFC7CE", "Needs Review": "FFEB9C"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance Matrix"
    if decision:
        ws.append([f"Decision: {decision}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
    ws.append([c.title() for c in COLUMNS])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS])
        fill = fills.get(r.get("status", ""))
        if fill:
            ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor=fill)
    widths = [8, 14, 10, 12, 24, 46, 34, 7, 9, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
